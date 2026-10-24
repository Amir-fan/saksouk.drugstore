from __future__ import annotations

import hashlib
from io import BytesIO
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageDraw, ImageFilter


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "public" / "medecines"
OUTPUT_DIR = ROOT / "public" / "data"
AUDIT_PATH = ROOT / "content" / "product-data-audit.json"
IMAGE_DIR = ROOT / "public" / "images" / "products"
BRAND_DIR = ROOT / "public" / "images" / "brands"
sys.stdout.reconfigure(encoding="utf-8")

COMPANIES = {
    "أدفانس": ("Advance", "أدفانس", "advance", "companies (2).jpeg"),
    "المشرق": ("Almashreq", "المشرق", "almashreq", ""),
    "ابن حيان": ("Ibn Hayan", "ابن حيان", "ibn-hayan", "companies (9).jpeg"),
    "انفليس": ("Infless", "إنفليس", "infless", "companies (1).png"),
    "بيوكسين": ("Bioxcin", "بيوكسين", "bioxcin", "companies (7).jpeg"),
    "بيوميد": ("Biomed", "بيوميد", "biomed", ""),
    "تانيوم": ("Tanium", "تانيوم", "tanium", "companies (8).jpeg"),
    "راشا": ("Rasha", "راشا", "rasha", "companies (3).png"),
    "زاروزا": ("Zaroza", "زاروزا", "zaroza", "companies (5).png"),
    "زين فارما": ("Zein Pharma", "زين فارما", "zein-pharma", "companies (6).jpeg"),
    "حماه فارما": ("Hama Pharma", "حماة فارما", "hama-pharma", "companies (3).jpeg"),
    "ميديكو": ("Medico", "ميديكو", "medico", "companies (1).jpeg"),
    "شفا": ("Shifa", "شفا", "shifa", "companies (10).jpeg"),
    "كيمي": ("Kimi", "كيمي", "kimi", "companies (4).png"),
}

COSMETIC_COMPANIES = {"Bioxcin", "Tanium", "Zaroza"}
SUPPLEMENT_BIOXCIN_MARKERS = ("TABLET", "SACHET", "5000MCG", "حبوب للشعر")
ADVANCE_COSMETIC_MARKERS = ("اوزوكير", "ليلي دوش")

# Supplied WhatsApp photos are deliberately mapped to catalogue names instead of
# relying on OCR. This keeps the import deterministic and makes repeated views a
# gallery on one product rather than duplicate products.
SUPPLIED_IMAGE_GROUPS: list[tuple[str, str, list[str]]] = [
    ("Advance", "بنتازايم", ["12.10.54 PM.jpeg", "12.10.57 PM (1).jpeg"]),
    ("Advance", "ادفا اوفاريا", ["12.10.55 PM (1).jpeg", "12.10.57 PM (3).jpeg"]),
    ("Advance", "ادفا غار", ["12.10.55 PM (2).jpeg", "12.10.55 PM.jpeg"]),
    ("Advance", "اوزوكير H PLUS", ["12.10.55 PM (3).jpeg"]),
    ("Advance", "ادفانسيد N.A.C", ["12.10.56 PM (1).jpeg"]),
    ("Advance", "اوزوكير M PLUS", ["12.10.56 PM (2).jpeg"]),
    ("Advance", "اوزوكير D PLUS", ["12.10.56 PM (3).jpeg"]),
    ("Advance", "ادفا مغنزيوم بلس", ["12.10.56 PM (4).jpeg"]),
    ("Advance", "الجيا كير", ["12.10.56 PM.jpeg"]),
    ("Advance", "اوزوكير R PLUS", ["12.10.57 PM (2).jpeg"]),
    ("Advance", "فومي رايت", ["12.10.57 PM (4).jpeg"]),
    ("Advance", "فيرو كير شراب", ["12.10.57 PM.jpeg"]),
    ("Bioxcin", "AQUA THERMAL ANTI DANDRUFF SHAMPOO", ["12.13.10 PM (1).jpeg"]),
    ("Bioxcin", "ATOPICARE NOURISHING FACE CREAM", ["12.13.10 PM (2).jpeg", "12.13.13 PM (4).jpeg"]),
    ("Bioxcin", "بلسم سائل", ["12.13.10 PM (3).jpeg", "12.13.11 PM.jpeg"]),
    ("Bioxcin", "FORTE SERUM", ["12.13.10 PM.jpeg", "12.13.14 PM (2).jpeg"]),
    ("Bioxcin", "BIOTIN SHAMPOO", ["12.13.11 PM (1).jpeg"]),
    ("Bioxcin", "COLLAGEN-BIOTIN LIQUID HAIR CONDITIONER", ["12.13.11 PM (2).jpeg", "12.13.11 PM (3).jpeg"]),
    ("Bioxcin", "COLD SKIN FACE CREAM", ["12.13.12 PM (1).jpeg", "12.13.17 PM (1).jpeg"]),
    ("Bioxcin", "QUANTUM OILY HAIR SHAMPOO", ["12.13.12 PM (2).jpeg"]),
    ("Bioxcin", "QUANTUM DRY/NORMAL HAIR SHAMPOO", ["12.13.12 PM (3).jpeg"]),
    ("Bioxcin", "COLLAGEN-BIOTIN HAIR CARE CREAM", ["12.13.12 PM.jpeg"]),
    ("Bioxcin", "COLLAGEN-BIOTIN HAIR SERUM", ["12.13.13 PM (1).jpeg"]),
    ("Bioxcin", "COLLAGEN-BIOTIN SHAMPOO", ["12.13.13 PM (2).jpeg", "12.13.14 PM.jpeg"]),
    ("Bioxcin", "KERATIN-ARGAN LIQUID CONDITIONER", ["12.13.13 PM (3).jpeg"]),
    ("Bioxcin", "WOMEN'S SKİN TABLET", ["12.13.13 PM.jpeg"]),
    ("Bioxcin", "KERATIN COMPLEX TABLET", ["12.13.14 PM (1).jpeg"]),
    ("Bioxcin", "FORTE SHAMPOO", ["12.13.15 PM (1).jpeg"]),
    ("Bioxcin", "MEN+SPORT SHAMPOO", ["12.13.15 PM (2).jpeg"]),
    ("Bioxcin", "DRY/NORMAL HAIR SHAMPOO", ["12.13.15 PM (3).jpeg"]),
    ("Bioxcin", "BLACK GARLIC SHAMPOO", ["12.13.15 PM.jpeg"]),
    ("Bioxcin", "KERATIN-ARGAN CARE OIL", ["12.13.16 PM (1).jpeg"]),
    ("Bioxcin", "QUANTUM SERUM", ["12.13.16 PM (2).jpeg"]),
    ("Bioxcin", "NOURISHING OILS SHAMPOO", ["12.13.16 PM (3).jpeg"]),
    ("Bioxcin", "KERATIN-ARGAN SHAMPOO", ["12.13.16 PM.jpeg"]),
    ("Bioxcin", "KERATIN-ARGAN HAIR CARE CREAM", ["12.13.17 PM (2).jpeg"]),
    ("Bioxcin", "DS AQUA THERMAL ANTI DANDRUFF SHAMPOO", ["12.13.17 PM.jpeg"]),
    ("Tanium", "Procapil", ["2.35.41 PM (1).jpeg"]),
    ("Tanium", "Avokado", ["2.35.41 PM (2).jpeg"]),
    ("Tanium", "Retinol", ["2.35.41 PM.jpeg"]),
    ("Tanium", "Glutathione", ["2.35.42 PM (1).jpeg"]),
    ("Tanium", "Caffein", ["2.35.42 PM (2).jpeg"]),
    ("Tanium", "Protein", ["2.35.42 PM (3).jpeg"]),
    ("Tanium", "Arbutin", ["2.35.42 PM.jpeg"]),
    ("Tanium", "Acne Care", ["2.35.43 PM (1).jpeg"]),
    ("Tanium", "Intensive تانيوم انتينسيف Mouth Wash", ["2.35.43 PM (2).jpeg"]),
    ("Tanium", "White Tinted", ["2.35.43 PM (3).jpeg"]),
    ("Tanium", "Hydra Fresh", ["2.35.43 PM.jpeg"]),
    ("Tanium", "Collagen", ["2.35.44 PM (1).jpeg"]),
    ("Tanium", "Intensive تانيوم انتينسيف معجون", ["2.35.44 PM (2).jpeg"]),
    ("Tanium", "Niacinamide", ["2.35.44 PM (3).jpeg"]),
    ("Tanium", "Sure Fresh", ["2.35.44 PM (4).jpeg"]),
    ("Tanium", "Pink Tinted", ["2.35.44 PM.jpeg"]),
    ("Tanium", "Hyaluronic", ["2.35.45 PM (1).jpeg"]),
    ("Tanium", "Beige Tinted", ["2.35.45 PM (2).jpeg"]),
    ("Tanium", "AHA-BHA", ["2.35.45 PM (3).jpeg"]),
    ("Tanium", "Light Skin", ["2.35.45 PM.jpeg"]),
    ("Tanium", "Gingivel تانيوم جينجيفال معجون", ["2.35.46 PM (1).jpeg"]),
    ("Tanium", "Sure Whitning", ["2.35.46 PM (2).jpeg"]),
    ("Tanium", "Vitamin-C", ["2.35.46 PM (3).jpeg"]),
    ("Tanium", "Gingivel تانيوم جينجيفال Mouth Wash", ["2.35.46 PM.jpeg"]),
]
ARABIC_RE = re.compile(r"[\u0600-\u06ff]")
ZERO_WIDTH_RE = re.compile(r"[\u200b-\u200f\u202a-\u202e\u2060-\u206f\ufeff]")


def clean(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = ZERO_WIDTH_RE.sub(" ", text)
    text = text.replace("\n", " ").replace("\r", " ")
    return " ".join(text.split()).strip(" -–—|،,")


def key_text(value: str) -> str:
    value = clean(value).casefold()
    value = value.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    return re.sub(r"[^0-9a-z\u0600-\u06ff]+", "", value)


def company_for(path: Path) -> tuple[str, str, str, str]:
    normalized = clean(path.stem)
    for marker, company in COMPANIES.items():
        if marker in normalized:
            return company
    raise ValueError(f"No company mapping for {path.name}")


def category_for(company_en: str, name: str) -> str:
    upper_name = clean(name).upper()
    if company_en == "Infless":
        return "supplements"
    if company_en == "Advance":
        return "cosmetics" if any(marker in clean(name) for marker in ADVANCE_COSMETIC_MARKERS) else "supplements"
    if company_en == "Bioxcin":
        return "supplements" if any(marker in upper_name for marker in SUPPLEMENT_BIOXCIN_MARKERS) else "cosmetics"
    if company_en == "Zaroza" and "HAIR VITAMIN" in upper_name:
        return "supplements"
    if company_en in COSMETIC_COMPANIES:
        return "cosmetics"
    return "medicine"


def canonical_name(company_en: str, name: str) -> str:
    normalized = clean(name)
    upper_name = normalized.upper()
    if company_en == "Bioxcin" and upper_name == "BIOXCIN FORTE SERUM 1X50":
        return "BIOXCIN FORTE SERUM"
    if company_en == "Zaroza" and "WASHING FOAM" in upper_name:
        return "ZAROZA FACE CLEANSING FOAM 150 ml"
    if company_en == "Tanium":
        if "(بيج" in normalized:
            return "Tanium Sun SPF 50+ — Beige Tinted تانيوم سن بيج"
        if normalized == "Tanium Sun تانيوم سن":
            return "Tanium Sun SPF 50+ — Pink Tinted تانيوم سن وردي"
        if normalized == "Tanium Sun":
            return "Tanium Sun SPF 50+ — White Tinted تانيوم سن أبيض"
    return normalized


def latin_prefix(value: str) -> str:
    match = ARABIC_RE.search(value)
    if not match:
        return clean(value)
    prefix = clean(value[: match.start()])
    return prefix if len(prefix) >= 2 else clean(value)


def make_product(
    path: Path,
    raw_name: str,
    *,
    display_name: str = "",
    description: str = "",
    composition: str = "",
    indication: str = "",
) -> dict[str, Any] | None:
    raw_name = clean(raw_name)
    if not raw_name:
        return None
    company_en, company_ar, company_slug, company_logo = company_for(path)
    category = category_for(company_en, raw_name)
    concise_name = clean(display_name) or (latin_prefix(raw_name) if category != "medicine" else raw_name)
    concise_name = concise_name or raw_name
    concise_name = canonical_name(company_en, concise_name)
    identity = f"{category}|{company_en}|{key_text(concise_name)}"
    return {
        "id": hashlib.sha1(identity.encode("utf-8")).hexdigest()[:12],
        "name": concise_name,
        "officialName": raw_name,
        "company": {
            "en": company_en,
            "ar": company_ar,
            "slug": company_slug,
            **({"logo": f"/images/brands/{company_slug}.webp"} if company_logo else {}),
        },
        "category": category,
        "group": "national" if category == "medicine" else category,
        "description": clean(description),
        "composition": clean(composition),
        "indication": clean(indication),
        "source": path.name,
    }


def image_map_for_sheet(sheet: Any) -> dict[int, bytes]:
    images: dict[int, bytes] = {}
    for image in getattr(sheet, "_images", []):
        anchor = getattr(image, "anchor", None)
        start = getattr(anchor, "_from", None)
        if start is None:
            continue
        try:
            images.setdefault(start.row + 1, image._data())
        except Exception:
            continue
    return images


def attach_image(product: dict[str, Any] | None, image_bytes: bytes | None) -> None:
    if not product or not image_bytes:
        return
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    relative_path = f"/images/products/catalog-{product['id']}.webp"
    output_path = ROOT / "public" / relative_path.lstrip("/")
    try:
        with Image.open(BytesIO(image_bytes)) as source:
            source.thumbnail((720, 720), Image.Resampling.LANCZOS)
            if source.mode not in {"RGB", "RGBA"}:
                source = source.convert("RGBA" if "transparency" in source.info else "RGB")
            source.save(output_path, "WEBP", quality=84, method=6)
        product["image"] = relative_path
        product["images"] = [relative_path]
    except Exception:
        return


def longest_true_run(flags: list[bool]) -> tuple[int, int] | None:
    """Return the longest meaningful run, bridging tiny compression gaps."""
    best: tuple[int, int] | None = None
    start: int | None = None
    last_true: int | None = None
    for index, flag in enumerate(flags):
        if flag:
            if start is None:
                start = index
            last_true = index
        elif start is not None and last_true is not None and index - last_true > 3:
            candidate = (start, last_true + 1)
            if best is None or candidate[1] - candidate[0] > best[1] - best[0]:
                best = candidate
            start = None
            last_true = None
    if start is not None and last_true is not None:
        candidate = (start, last_true + 1)
        if best is None or candidate[1] - candidate[0] > best[1] - best[0]:
            best = candidate
    return best


def trim_capture_bands(source: Image.Image) -> Image.Image:
    """Remove black phone/gallery chrome without cropping normal product photos.

    Supplied product imagery sometimes arrives as a screenshot whose actual photo
    occupies the middle of a black canvas. We find the longest region containing
    real image information and only crop when a substantial outer band is present.
    """
    image = source.convert("RGB")
    width, height = image.size
    if width < 80 or height < 80:
        return image

    preview = image.copy()
    preview.thumbnail((260, 260), Image.Resampling.BILINEAR)
    sample_width, sample_height = preview.size
    pixels = preview.load()

    def visible(pixel: tuple[int, int, int]) -> bool:
        return max(pixel) > 28

    row_flags = [
        sum(1 for x in range(sample_width) if visible(pixels[x, y])) / sample_width >= .12
        for y in range(sample_height)
    ]
    row_run = longest_true_run(row_flags)
    if row_run is None or row_run[1] - row_run[0] < sample_height * .18:
        return image

    top = max(0, round(row_run[0] / sample_height * height) - round(height * .004))
    bottom = min(height, round(row_run[1] / sample_height * height) + round(height * .004))
    vertical_crop = image.crop((0, top, width, bottom))

    column_preview = vertical_crop.copy()
    column_preview.thumbnail((260, 260), Image.Resampling.BILINEAR)
    column_width, column_height = column_preview.size
    column_pixels = column_preview.load()
    column_flags = [
        sum(1 for y in range(column_height) if visible(column_pixels[x, y])) / column_height >= .12
        for x in range(column_width)
    ]
    column_run = longest_true_run(column_flags)
    if column_run is None or column_run[1] - column_run[0] < column_width * .18:
        return vertical_crop if vertical_crop.height < height * .94 else image

    left = max(0, round(column_run[0] / column_width * width) - round(width * .004))
    right = min(width, round(column_run[1] / column_width * width) + round(width * .004))
    cropped = image.crop((left, top, right, bottom))

    removed_fraction = 1 - (cropped.width * cropped.height) / (width * height)
    return cropped if removed_fraction >= .06 else image


def remove_connected_dark_background(source: Image.Image) -> Image.Image:
    """Turn a solid black outer canvas transparent while retaining dark label art."""
    image = source.convert("RGBA")
    width, height = image.size
    preview = image.convert("RGB")
    preview.thumbnail((180, 180), Image.Resampling.BILINEAR)
    preview_width, preview_height = preview.size
    preview_pixels = preview.load()
    perimeter = [preview_pixels[x, 0] for x in range(preview_width)]
    perimeter += [preview_pixels[x, preview_height - 1] for x in range(preview_width)]
    perimeter += [preview_pixels[0, y] for y in range(preview_height)]
    perimeter += [preview_pixels[preview_width - 1, y] for y in range(preview_height)]
    if sum(1 for pixel in perimeter if max(pixel) < 55) / len(perimeter) < .18:
        return image

    rgb = image.convert("RGB")
    dark = Image.eval(rgb.convert("L"), lambda value: 255 if value < 48 else 0)
    # A one-pixel dark border joins every outside-connected component, allowing
    # one flood fill instead of scanning/filling thousands of perimeter seeds.
    padded = Image.new("L", (width + 2, height + 2), 255)
    padded.paste(dark, (1, 1))
    ImageDraw.floodfill(padded, (0, 0), 128)
    connected_background = padded.crop((1, 1, width + 1, height + 1))
    connected_background = connected_background.point(lambda value: 255 if value == 128 else 0)
    connected_background = connected_background.filter(ImageFilter.GaussianBlur(radius=1.1))
    image.putalpha(ImageChops.invert(connected_background))
    return image


def optimize_supplied_image(source_path: Path, product: dict[str, Any], index: int) -> str | None:
    output_name = f"product-{product['id']}-{index + 1}.webp"
    output_path = IMAGE_DIR / output_name
    try:
        with Image.open(source_path) as source:
            source = trim_capture_bands(source)
            source = remove_connected_dark_background(source)
            source.thumbnail((1200, 1200), Image.Resampling.LANCZOS)
            source.save(output_path, "WEBP", quality=88, method=6)
        return f"/images/products/{output_name}"
    except Exception:
        return None


def supplied_image_path(suffix: str) -> Path | None:
    matches = sorted(IMAGE_DIR.glob(f"WhatsApp Image 2026-10-* at {suffix}"))
    return matches[0] if matches else None


def attach_supplied_images(products: list[dict[str, Any]]) -> list[dict[str, str]]:
    exact_hashes: set[str] = set()
    audit: list[dict[str, str]] = []
    for company, product_marker, suffixes in SUPPLIED_IMAGE_GROUPS:
        marker_key = key_text(product_marker)
        candidates = [
            item
            for item in products
            if item["company"]["en"] == company
            and marker_key in key_text(f"{item['name']} {item.get('officialName', '')}")
        ]
        product = min(
            candidates,
            key=lambda item: len(key_text(f"{item['name']} {item.get('officialName', '')}")) - len(marker_key),
            default=None,
        )
        if product is None:
            audit.append({"company": company, "product": product_marker, "status": "product-not-found"})
            continue
        supplied: list[str] = []
        for suffix in suffixes:
            source = supplied_image_path(suffix)
            if source is None:
                existing_name = f"product-{product['id']}-{len(supplied) + 1}.webp"
                if (IMAGE_DIR / existing_name).exists():
                    supplied.append(f"/images/products/{existing_name}")
                    audit.append({"company": company, "product": product_marker, "status": f"using-optimized: {existing_name}"})
                else:
                    audit.append({"company": company, "product": product_marker, "status": f"image-not-found: {suffix}"})
                continue
            digest = hashlib.sha256(source.read_bytes()).hexdigest()
            if digest in exact_hashes:
                audit.append({"company": company, "product": product_marker, "status": f"exact-duplicate-skipped: {source.name}"})
                continue
            exact_hashes.add(digest)
            optimized = optimize_supplied_image(source, product, len(supplied))
            if optimized:
                supplied.append(optimized)
        if supplied:
            embedded = [path for path in product.get("images", []) if path not in supplied]
            product["images"] = supplied + embedded
            product["image"] = supplied[0]
            audit.append({"company": company, "product": product["name"], "status": f"matched-{len(supplied)}"})
    return audit


def build_brand_logos() -> None:
    BRAND_DIR.mkdir(parents=True, exist_ok=True)
    source_dir = ROOT / "public" / "images"
    for _, (_, _, slug, source_name) in COMPANIES.items():
        if not source_name:
            continue
        source_path = source_dir / source_name
        output_path = BRAND_DIR / f"{slug}.webp"
        with Image.open(source_path) as source:
            source = source.convert("RGBA")
            source.thumbnail((640, 260), Image.Resampling.LANCZOS)
            background = Image.new("RGBA", source.size, "white")
            background.alpha_composite(source)
            background.convert("RGB").save(output_path, "WEBP", quality=92, method=6)


def first_sheet_rows(path: Path) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    return list(workbook.worksheets[0].iter_rows(values_only=True))


def extract_advance(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook.worksheets[0]
    images = image_map_for_sheet(sheet)
    products = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number < 3:
            continue
        name = clean(row[1] if len(row) > 1 else "")
        if not name:
            continue
        product = make_product(
            path,
            name,
            composition=row[2] if len(row) > 2 else "",
            indication=row[3] if len(row) > 3 else "",
        )
        if product:
            attach_image(product, images.get(row_number))
            products.append(product)
    return products


def extract_described_cosmetics(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook.worksheets[0]
    images = image_map_for_sheet(sheet)
    products = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        value = clean(row[0] if row else "")
        if not value or value in {"اسم المنتج", "الصنف"}:
            continue
        product = make_product(path, value, description=value)
        if product:
            attach_image(product, images.get(row_number))
            products.append(product)
    return products


def extract_zaroza(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    products = []
    for sheet in workbook.worksheets:
        images = image_map_for_sheet(sheet)
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [clean(value) for value in row]
            if not any(values):
                continue
            candidate = values[2] if len(values) > 2 else values[-1]
            numbered = bool(values and re.fullmatch(r"\d+", values[0]))
            if not numbered and sheet.title != "Table 2":
                continue
            if not candidate:
                continue
            product = make_product(path, candidate, description=candidate)
            if product:
                attach_image(product, images.get(row_number))
                products.append(product)
    return products


def extract_tanium(path: Path) -> list[dict[str, Any]]:
    document = Document(path)
    products = []
    for table in document.tables:
        for row in table.rows[1:]:
            value = clean(row.cells[0].text)
            if not value:
                continue
            product = make_product(path, value, display_name=value, description=value)
            if product:
                image_bytes = None
                for blip in row._tr.xpath(".//a:blip"):
                    relationship_id = blip.get(qn("r:embed"))
                    related_part = document.part.related_parts.get(relationship_id)
                    if related_part is not None and hasattr(related_part, "blob"):
                        image_bytes = related_part.blob
                        break
                attach_image(product, image_bytes)
                products.append(product)
    return products


def extract_medico(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    products = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        name_column = None
        composition_column = None
        for row in rows[:15]:
            for index, value in enumerate(row):
                text = clean(value)
                if "اسم المستحضر" in text or "اسم الصنف" in text:
                    name_column = index
                if "التركيب والعيار" in text:
                    composition_column = index
        if name_column is None:
            nonempty_columns = Counter(
                index
                for row in rows[:15]
                for index, value in enumerate(row)
                if clean(value)
            )
            if not nonempty_columns:
                continue
            name_column = max(nonempty_columns)
        for row in rows:
            if name_column >= len(row):
                continue
            name = clean(row[name_column])
            if not name:
                continue
            if any(
                phrase in name
                for phrase in (
                    "اسم المستحضر",
                    "اسم الصنف",
                    "تاريخ الترخيص",
                    "لائحة بالأصناف",
                )
            ):
                continue
            nonempty = [clean(value) for value in row if clean(value)]
            if len(nonempty) < 2:
                continue
            composition = ""
            if composition_column is not None and composition_column < len(row):
                composition = clean(row[composition_column])
            product = make_product(path, name, composition=composition)
            if product:
                products.append(product)
    return products


def extract_simple(path: Path) -> list[dict[str, Any]]:
    rows = first_sheet_rows(path)
    start = 0
    stem = clean(path.stem)
    if "زين فارما" in stem:
        start = 5
    elif "حماه فارما" in stem:
        start = 2
    products = []
    headers = {"Issue No.", "Issue Date", "اسم الصنف", "الصنف", "اسم المنتج"}
    for row in rows[start:]:
        name = clean(row[0] if row else "")
        if not name or name in headers:
            continue
        product = make_product(path, name)
        if product:
            products.append(product)
    return products


def extract(path: Path) -> list[dict[str, Any]]:
    stem = clean(path.stem)
    if "أدفانس" in stem:
        return extract_advance(path)
    if path.suffix.lower() == ".docx":
        return extract_tanium(path)
    if "ميديكو" in stem:
        return extract_medico(path)
    if "زاروزا" in stem:
        return extract_zaroza(path)
    if "كوزمتك" in stem:
        return extract_described_cosmetics(path)
    return extract_simple(path)


all_products: list[dict[str, Any]] = []
source_audit: list[dict[str, Any]] = []
IMAGE_DIR.mkdir(parents=True, exist_ok=True)
for old_image in IMAGE_DIR.glob("catalog-*.webp"):
    old_image.unlink()
if any(IMAGE_DIR.glob("WhatsApp Image 2026-10-*.jpeg")):
    for old_image in IMAGE_DIR.glob("product-*.webp"):
        old_image.unlink()
for source_path in sorted(SOURCE_DIR.iterdir(), key=lambda item: item.name):
    if source_path.suffix.lower() not in {".xlsx", ".docx"}:
        continue
    extracted = extract(source_path)
    all_products.extend(extracted)
    source_audit.append(
        {
            "source": source_path.name,
            "categoryBreakdown": dict(Counter(item["category"] for item in extracted)),
            "extractedRows": len(extracted),
        }
    )

# Two photographed products were not listed as text rows in the supplied files.
# Adding them here keeps the images and catalogue connected without inventing
# formulations or claims that were not provided.
advance_source = next(path for path in SOURCE_DIR.iterdir() if "أدفانس" in path.stem)
tanium_source = next(path for path in SOURCE_DIR.iterdir() if "تانيوم" in path.stem)
all_products.extend(
    product
    for product in (
        make_product(advance_source, "اوزوكير D PLUS"),
        make_product(tanium_source, "Tanium Procapil Hair Serum سيروم بروكابيل للشعر"),
    )
    if product
)

deduplicated: list[dict[str, Any]] = []
duplicates: list[dict[str, str]] = []
seen: dict[str, dict[str, Any]] = {}
for product in all_products:
    duplicate_key = f"{product['category']}|{product['company']['en']}|{key_text(product['name'])}"
    if duplicate_key in seen:
        kept = seen[duplicate_key]
        for image in product.get("images", []):
            kept.setdefault("images", [])
            if image not in kept["images"]:
                kept["images"].append(image)
        for field in ("description", "composition", "indication"):
            if not kept.get(field) and product.get(field):
                kept[field] = product[field]
        duplicates.append(
            {
                "source": product["source"],
                "name": product["name"],
                "keptId": seen[duplicate_key]["id"],
            }
        )
        continue
    seen[duplicate_key] = product
    deduplicated.append(product)

image_matches = attach_supplied_images(deduplicated)
build_brand_logos()

for product in deduplicated:
    product.pop("source", None)
    if product["officialName"] == product["name"]:
        product["officialName"] = ""
    for field in ("officialName", "description", "composition", "indication"):
        if not product[field]:
            product.pop(field)
    if product.get("images"):
        product["images"] = list(dict.fromkeys(product["images"]))

referenced_images = {product.get("image") for product in deduplicated if product.get("image")}
referenced_images.update(
    image for product in deduplicated for image in product.get("images", [])
)
for generated_image in IMAGE_DIR.glob("catalog-*.webp"):
    relative_image = f"/images/products/{generated_image.name}"
    if relative_image not in referenced_images:
        generated_image.unlink()
for generated_image in IMAGE_DIR.glob("product-*.webp"):
    relative_image = f"/images/products/{generated_image.name}"
    if relative_image not in referenced_images:
        generated_image.unlink()

deduplicated.sort(key=lambda product: (product["category"], product["company"]["en"], product["name"].casefold()))
catalogue = {
    "medicine": [product for product in deduplicated if product["category"] == "medicine"],
    "cosmetics": [product for product in deduplicated if product["category"] == "cosmetics"],
    "supplements": [product for product in deduplicated if product["category"] == "supplements"],
}
audit = {
    "sources": source_audit,
    "rawExtractedRows": len(all_products),
    "duplicateRowsRemoved": len(duplicates),
    "uniqueProducts": len(deduplicated),
    "categoryCounts": {category: len(products) for category, products in catalogue.items()},
    "productsWithImages": sum(1 for product in deduplicated if product.get("image")),
    "duplicates": duplicates,
    "suppliedImageMatches": image_matches,
}

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
for category, products in catalogue.items():
    (OUTPUT_DIR / f"products-{category}.json").write_text(
        json.dumps(products, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
AUDIT_PATH.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({key: value for key, value in audit.items() if key != "duplicates"}, ensure_ascii=False, indent=2))
